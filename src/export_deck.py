"""Render the presentation: PowerPoint with speaker notes + a standalone script.

    python -m src.export_deck

  slides/earnings_discontinuity.pptx   16:9 deck, notes in every slide's notes pane
  slides/speaker_notes.md              the same script, one section per slide

Drag the .pptx into Google Drive to open it as Google Slides (notes carry over).
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .deck_content import AFFIL, AUTHOR, SLIDES

ROOT = Path(__file__).resolve().parents[1]
OUT_PPTX = ROOT / "slides" / "earnings_discontinuity.pptx"
OUT_MD = ROOT / "slides" / "speaker_notes.md"

W, H = Inches(13.333), Inches(7.5)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x84, 0x83, 0x7C)
ACCENT = RGBColor(0x2A, 0x78, 0xD6)
ACCENT2 = RGBColor(0xEB, 0x68, 0x34)
RULE = RGBColor(0xE3, 0xE2, 0xDD)
SURFACE = RGBColor(0xFC, 0xFC, 0xFB)
FONT = "Calibri"

MARGIN = Inches(0.6)


def _bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = SURFACE


def _text(slide, left, top, width, height, text, *, size, bold=False, color=INK,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def _kicker_title(slide, kicker: str, title: str) -> Emu:
    _text(slide, MARGIN, Inches(0.42), W - 2 * MARGIN, Inches(0.3), kicker,
          size=12, bold=True, color=ACCENT)
    size = 28 if len(title) <= 70 else 24
    box = _text(slide, MARGIN, Inches(0.72), W - 2 * MARGIN, Inches(1.0), title,
                size=size, bold=True, color=INK, line_spacing=1.05)
    lines = 1 if len(title) <= 80 else 2
    return Inches(0.72) + Inches(0.5 * lines) + Inches(0.25)


def _picture_fit(slide, path: Path, left, top, width, height):
    with Image.open(path) as im:
        w, h = im.size
    scale = min(width / w, height / h)
    pw, ph = int(w * scale), int(h * scale)
    pl = left + (width - pw) // 2
    pt = top + (height - ph) // 2
    return slide.shapes.add_picture(str(path), pl, pt, pw, ph)


def _bullets(slide, items, left, top, width, height, *, head_size=17, body_size=14):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    first = True
    for head, body in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(0) if p is tf.paragraphs[0] else Pt(12)
        p.line_spacing = 1.08
        r = p.add_run()
        r.text = head
        r.font.name = FONT
        r.font.size = Pt(head_size)
        r.font.bold = True
        r.font.color.rgb = INK
        p2 = tf.add_paragraph()
        p2.line_spacing = 1.12
        p2.space_before = Pt(2)
        r2 = p2.add_run()
        r2.text = body
        r2.font.name = FONT
        r2.font.size = Pt(body_size)
        r2.font.color.rgb = INK2
    return box


def _table(slide, columns, rows, left, top, width, *, size=12, col_widths=None):
    n_rows, n_cols = len(rows) + 1, len(columns)
    row_h = Inches(0.36)
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, row_h * n_rows)
    tbl = shape.table
    if col_widths:
        total = sum(col_widths)
        for c, cw in enumerate(col_widths):
            tbl.columns[c].width = int(width * cw / total)
    z_cols = {c for c, name in enumerate(columns) if name.strip().lower().startswith("z") or name.strip().lower().endswith(" z")}
    for c, name in enumerate(columns):
        cell = tbl.cell(0, c)
        cell.text = name
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0xF3, 0xF2, 0xEE)
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
            for r in p.runs:
                r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), True, INK
    for ri, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.cell(ri, c)
            txt = str(val)
            cell.text = txt
            cell.fill.solid()
            cell.fill.fore_color.rgb = SURFACE if ri % 2 else RGBColor(0xF8, 0xF7, 0xF4)
            hit = False
            if c in z_cols:
                try:
                    hit = abs(float(txt.replace("\u2212", "-").split(" ")[0])) > 1.96
                except ValueError:
                    hit = False
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
                for r in p.runs:
                    r.font.name, r.font.size = FONT, Pt(size)
                    r.font.color.rgb = ACCENT2 if hit else INK2
                    r.font.bold = hit
    return shape


def _footer(slide, n: int, total: int) -> None:
    _text(slide, MARGIN, H - Inches(0.42), Inches(8), Inches(0.25),
          "Earnings discontinuity at zero · SEC 10-K filings, fiscal years 2016–2025",
          size=9, color=MUTED)
    _text(slide, W - MARGIN - Inches(1.2), H - Inches(0.42), Inches(1.2), Inches(0.25),
          f"{n} / {total}", size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def build() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    total = len(SLIDES)

    for n, spec in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        _bg(slide)
        kind = spec["kind"]

        if kind == "title":
            _text(slide, MARGIN, Inches(2.2), W - 2 * MARGIN, Inches(1.8), spec["title"],
                  size=40, bold=True, color=INK, line_spacing=1.05)
            _text(slide, MARGIN, Inches(4.15), W - 2 * MARGIN - Inches(2), Inches(1.0), spec["subtitle"],
                  size=18, color=INK2, line_spacing=1.2)
            _text(slide, MARGIN, Inches(6.1), Inches(8), Inches(0.35), AUTHOR, size=14, bold=True, color=INK)
            _text(slide, MARGIN, Inches(6.45), Inches(8), Inches(0.35), AFFIL, size=12, color=MUTED)
            line = slide.shapes.add_shape(1, MARGIN, Inches(5.85), Inches(1.2), Inches(0.05))
            line.fill.solid(); line.fill.fore_color.rgb = ACCENT; line.line.fill.background()

        elif kind == "image":
            top = _kicker_title(slide, spec["kicker"], spec["title"])
            _picture_fit(slide, ROOT / spec["image"], MARGIN, top, W - 2 * MARGIN, H - top - Inches(0.6))
            _footer(slide, n, total)

        elif kind == "bullets":
            top = _kicker_title(slide, spec["kicker"], spec["title"])
            n_items = len(spec["bullets"])
            head, body = (17, 14) if n_items <= 4 else (15, 12.5)
            _bullets(slide, spec["bullets"], MARGIN, top + Inches(0.1), W - 2 * MARGIN, H - top - Inches(0.8),
                     head_size=head, body_size=body)
            _footer(slide, n, total)

        elif kind == "table":
            top = _kicker_title(slide, spec["kicker"], spec["title"])
            n_cols = len(spec["columns"])
            size = spec.get("font", 13 if n_cols <= 7 else 11)
            shape = _table(slide, spec["columns"], spec["rows"], MARGIN, top + Inches(0.15),
                           W - 2 * MARGIN, size=size, col_widths=spec.get("col_widths"))
            cap_top = top + Inches(0.15) + Inches(0.36) * (len(spec["rows"]) + 1) + Inches(0.25)
            _text(slide, MARGIN, cap_top, W - 2 * MARGIN, Inches(0.8), spec["caption"],
                  size=11, color=MUTED, line_spacing=1.15)
            _footer(slide, n, total)

        elif kind == "image_table":
            top = _kicker_title(slide, spec["kicker"], spec["title"])
            img_w = Inches(7.0)
            _picture_fit(slide, ROOT / spec["image"], MARGIN, top, img_w, H - top - Inches(0.7))
            tleft = MARGIN + img_w + Inches(0.3)
            twidth = W - tleft - MARGIN
            _table(slide, spec["columns"], spec["rows"], tleft, top + Inches(0.3), twidth,
                   size=spec.get("font", 11), col_widths=spec.get("col_widths"))
            cap_top = top + Inches(0.3) + Inches(0.36) * (len(spec["rows"]) + 1) + Inches(0.3)
            _text(slide, tleft, cap_top, twidth, Inches(1.6), spec["caption"], size=11, color=MUTED, line_spacing=1.15)
            _footer(slide, n, total)

        slide.notes_slide.notes_text_frame.text = spec["notes"]

    OUT_PPTX.parent.mkdir(exist_ok=True)
    prs.save(OUT_PPTX)

    lines = ["# Speaker notes — Earnings discontinuity at zero", "",
             f"{AUTHOR} · {AFFIL}", "",
             "The same script that sits in the notes pane of `earnings_discontinuity.pptx`, one section per slide. "
             "Written to be said out loud. Roughly ten minutes for slides 1–15; 16–19 are backup.", ""]
    for n, spec in enumerate(SLIDES, start=1):
        head = spec.get("kicker", "TITLE")
        lines += [f"## Slide {n} — {spec['title']}", "", f"*{head}*", "", spec["notes"], ""]
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {OUT_PPTX.relative_to(ROOT)} ({total} slides, {OUT_PPTX.stat().st_size // 1024} KB) and {OUT_MD.relative_to(ROOT)}")


if __name__ == "__main__":
    build()
